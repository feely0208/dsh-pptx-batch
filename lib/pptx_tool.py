#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
dsh-pptx-batch 的真实处理逻辑（python-pptx）

由 bin/cli.mjs 通过 spawnSync 调用：
    python3 lib/pptx_tool.py <action> '<json payload>'

约定：
  - stdout 只输出一行 JSON（人类可读文案由 Node 侧负责）
  - stderr 输出异常栈（调试用）
  - 退出码：0 成功 / 1 运行失败 / 2 参数或输入问题 / 3 缺依赖
  - 只读输入文件，所有产物写进 --out 目录（默认 <输入目录>/_输出）
"""

import json
import os
import re
import shutil
import struct
import sys
import tempfile
import traceback
import zlib
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches
except ImportError as exc:  # 缺依赖必须明确提示安装命令，不能静默
    print(json.dumps({
        "ok": False,
        "code": 3,
        "message": "缺少 Python 依赖：%s" % getattr(exc, "name", exc),
        "hint": "请先安装：\n  python3 -m pip install python-pptx\n"
                "如果 python3 也没有：\n  macOS: brew install python3\n  Ubuntu: sudo apt install python3",
    }, ensure_ascii=False))
    sys.exit(3)

PPTX_EXT = ".pptx"
SKIP_DIRS = {"_输出", "node_modules", ".git", "__pycache__", ".venv", "venv"}
PLACEHOLDER_RE = re.compile(r"\{\{\s*([^{}]{1,60}?)\s*\}\}")
DEMO_TITLE = "季度经营简报"
DEMO_CLIENT = "示例客户甲"
DEMO_AMOUNT = "¥ 128,000"


# ── 错误 ──────────────────────────────────────────────────────────

class ToolError(Exception):
    def __init__(self, message, hint="", code=2):
        super().__init__(message)
        self.message = message
        self.hint = hint
        self.code = code


def emit(obj):
    print(json.dumps(obj, ensure_ascii=False))


# ── 路径与文件收集 ────────────────────────────────────────────────

def natural_key(text):
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", text)]


def resolve_out(payload, input_path):
    """输出目录：--out 优先，否则 <输入目录>/_输出。"""
    out = payload.get("out")
    if out:
        return Path(out).expanduser().resolve()
    base = Path(input_path).expanduser().resolve()
    if not base.is_dir():
        base = base.parent
    return base / "_输出"


def collect_pptx(input_path, out_dir=None):
    """接受单个 .pptx 或目录（递归），按文件名自然排序。"""
    if not input_path:
        raise ToolError("缺少 --input", "例：fill --input ./模板目录 --set 客户=客户甲")
    root = Path(input_path).expanduser().resolve()
    if not root.exists():
        raise ToolError(
            "输入不存在：%s" % input_path,
            "请检查路径是否正确；传目录会递归处理其中的 .pptx 文件。",
        )
    skip_out = set()
    if out_dir:
        try:
            skip_out = {str(Path(out_dir).resolve())}
        except OSError:
            skip_out = set()
    if root.is_file():
        if root.suffix.lower() != PPTX_EXT:
            raise ToolError(
                "输入不是 .pptx 文件：%s" % root.name,
                "本插件只处理 PowerPoint 2007+ 的 .pptx；.ppt 请先另存为 .pptx。",
            )
        return [root]
    files = []
    for dirpath, dirnames, filenames in os.walk(root):
        current = str(Path(dirpath).resolve())
        if current in skip_out:
            dirnames[:] = []
            continue
        dirnames[:] = sorted(
            [d for d in dirnames if d not in SKIP_DIRS and not d.startswith("~$")],
            key=natural_key,
        )
        for name in filenames:
            if name.startswith("~$"):
                continue  # PowerPoint 打开时的临时文件
            p = Path(dirpath) / name
            if p.suffix.lower() == PPTX_EXT:
                files.append(p)
    files.sort(key=lambda p: natural_key(str(p.relative_to(root))))
    if not files:
        raise ToolError(
            "目录里没有 .pptx 文件：%s" % root,
            "确认目录选对了；本插件不处理 .ppt（请先另存为 .pptx）。",
        )
    return files


def unique_path(directory, filename):
    """同名文件不覆盖，追加 _1/_2 序号。"""
    directory = Path(directory)
    target = directory / filename
    if not target.exists():
        return target
    stem, suffix = target.stem, target.suffix
    n = 1
    while (directory / ("%s_%d%s" % (stem, n, suffix))).exists():
        n += 1
    return directory / ("%s_%d%s" % (stem, n, suffix))


# ── 形状遍历 ──────────────────────────────────────────────────────

def shape_type_of(shape):
    try:
        return shape.shape_type
    except Exception:
        return None


def is_group(shape):
    return shape_type_of(shape) == MSO_SHAPE_TYPE.GROUP


def is_picture(shape):
    return shape_type_of(shape) == MSO_SHAPE_TYPE.PICTURE


def iter_shapes(shapes):
    """按文档顺序遍历形状，组合形状递归展开（只读遍历，不改变结构）。"""
    for shape in shapes:
        yield shape
        if is_group(shape):
            try:
                yield from iter_shapes(shape.shapes)
            except Exception:
                continue  # 极少数残缺组合形状取不到子形状：跳过


def shape_kind(shape):
    if getattr(shape, "has_table", False):
        return "table"
    if getattr(shape, "has_chart", False):
        return "chart"
    if is_picture(shape):
        return "picture"
    if is_group(shape):
        return "group"
    if getattr(shape, "has_text_frame", False):
        return "text"
    return "shape"


def shape_paragraphs(shape):
    """形状里可编辑的段落：文本框 + 表格单元格。"""
    if getattr(shape, "has_text_frame", False):
        yield from shape.text_frame.paragraphs
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                yield from cell.text_frame.paragraphs


def iter_paragraphs(prs):
    """整份 PPT 里所有可编辑段落（每页的文本框、表格单元格，组合形状递归）。"""
    for slide in prs.slides:
        for shape in iter_shapes(slide.shapes):
            yield from shape_paragraphs(shape)


# ── 文本替换（尽量保留 run 级格式） ───────────────────────────────

def sub_paragraph(paragraph, pattern, replacement):
    """
    在单个段落里按正则替换，返回替换次数。

    先按 run 拼出整段文字再定位匹配，然后把「匹配外的字符」原样归还原来的 run，
    「替换后的文字」落到匹配起点所在的 run 上 —— 所以没被匹配碰到的 run 格式完全不变，
    只有真正跨 run 的匹配才会把这几段文字并到起始 run 的格式上。
    """
    runs = list(paragraph.runs)
    if not runs:
        return 0
    texts = [run.text for run in runs]
    full = "".join(texts)
    if not full:
        return 0
    matches = list(pattern.finditer(full))
    if not matches:
        return 0
    owner = []
    for idx, text in enumerate(texts):
        owner.extend([idx] * len(text))
    pieces = [[] for _ in runs]
    cursor = 0
    for match in matches:
        start, end = match.span()
        if end <= cursor:
            continue
        if start > cursor:
            for pos in range(cursor, start):
                pieces[owner[pos]].append(full[pos])
        pieces[owner[start]].append(replacement)
        cursor = end
    for pos in range(cursor, len(full)):
        pieces[owner[pos]].append(full[pos])
    for run, old, parts in zip(runs, texts, pieces):
        new = "".join(parts)
        if new != old:
            run.text = new
    return len(matches)


def paragraph_text(paragraph):
    try:
        return paragraph.text
    except Exception:
        return ""


def leftover_placeholders(prs):
    """填完之后还剩哪些 {{...}} —— 帮用户抓模板里拼错的 key。"""
    left = []
    for paragraph in iter_paragraphs(prs):
        for match in PLACEHOLDER_RE.finditer(paragraph_text(paragraph)):
            key = match.group(1).strip()
            if key and key not in left:
                left.append(key)
    return left


# ── 各动作 ────────────────────────────────────────────────────────

def parse_sets(raw_sets):
    if not raw_sets:
        raise ToolError(
            "缺少 --set（要填的占位符）",
            "每个值用 key=value，可重复：\n"
            "  fill --input ./模板 --set 客户=客户甲 --set 金额=1000",
        )
    pairs, seen = [], set()
    for item in raw_sets:
        if "=" not in item:
            raise ToolError(
                "--set 需要 key=value 形式：%s" % item,
                "例：--set 客户=客户甲（等号左边是模板里的 {{key}}，右边是要填的值）",
            )
        key, value = item.split("=", 1)
        key = key.strip()
        if not key:
            raise ToolError(
                "--set 的 key 不能为空：%s" % item,
                "例：--set 客户=客户甲",
            )
        if key in seen:
            raise ToolError(
                "--set 重复给了同一个 key：%s" % key,
                "同一个 key 只写一次，需要不同值请分两次运行。",
            )
        seen.add(key)
        pairs.append((key, value))
    return pairs


def action_fill(payload):
    files = collect_pptx(payload.get("input"))
    out = resolve_out(payload, payload.get("input"))
    pairs = parse_sets(payload.get("set"))
    out.mkdir(parents=True, exist_ok=True)
    results, total, unfilled = [], 0, []
    for src in files:
        prs = Presentation(str(src))
        counts = {}
        for key, value in pairs:
            pattern = re.compile(r"\{\{\s*" + re.escape(key) + r"\s*\}\}")
            for paragraph in iter_paragraphs(prs):
                n = sub_paragraph(paragraph, pattern, value)
                if n:
                    counts[key] = counts.get(key, 0) + n
        file_total = sum(counts.values())
        left = leftover_placeholders(prs)
        for key in left:
            if key not in unfilled:
                unfilled.append(key)
        target = unique_path(out, src.name)
        prs.save(str(target))
        results.append({
            "input": str(src),
            "input_name": src.name,
            "output": str(target),
            "output_name": target.name,
            "replacements": file_total,
            "keys": counts,
            "unfilled": left,
        })
        total += file_total
    summary = "处理 %d 个文件，共替换 %d 处占位符" % (len(files), total)
    if unfilled:
        summary += "；仍有未赋值占位符 %d 个：%s" % (len(unfilled), "、".join(unfilled[:8]))
    return {
        "action": "fill",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "out": str(out),
        "files": results,
        "replacements": total,
        "unfilled": unfilled,
        "summary": summary,
        "artifacts": [r["output"] for r in results],
    }


def picture_names(prs):
    names = []
    for slide in prs.slides:
        for shape in iter_shapes(slide.shapes):
            if is_picture(shape) and shape.name and shape.name not in names:
                names.append(shape.name)
    return names


def action_replace_image(payload):
    files = collect_pptx(payload.get("input"))
    out = resolve_out(payload, payload.get("input"))
    shape_name = (payload.get("shape") or "").strip()
    image_arg = (payload.get("image") or "").strip()
    if not shape_name:
        raise ToolError(
            "缺少 --shape（要替换的图片形状名）",
            "例：replace-image --input ./模板 --shape Logo --image ./logo.png\n"
            "形状名要和 PowerPoint 选择窗格（或「开始 → 选择 → 选择窗格」）里显示的名字完全一致。",
        )
    if not image_arg:
        raise ToolError(
            "缺少 --image（新图片文件）",
            "例：replace-image --input ./模板 --shape Logo --image ./logo.png",
        )
    image = Path(image_arg).expanduser().resolve()
    if not image.exists() or not image.is_file():
        raise ToolError(
            "新图片不存在：%s" % image_arg,
            "确认 --image 指向一个 png / jpg / gif / bmp / tiff 图片文件。",
        )
    out.mkdir(parents=True, exist_ok=True)
    results, total = [], 0
    non_picture_hits, all_picture_names = [], []
    for src in files:
        prs = Presentation(str(src))
        replaced, non_picture = 0, []
        for slide in prs.slides:
            rId = None
            for shape in iter_shapes(slide.shapes):
                if shape.name != shape_name:
                    continue
                if not is_picture(shape):
                    non_picture.append(shape_kind(shape))
                    continue
                if rId is None:
                    try:
                        _part, rId = slide.part.get_or_add_image_part(str(image))
                    except Exception:
                        raise ToolError(
                            "新图片无法读取：%s" % image.name,
                            "支持 png / jpg / gif / bmp / tiff；确认文件没损坏、后缀与内容一致。",
                        )
                # 只换图片内容，保留原形状的位置与尺寸（版面不动）
                shape._element.blipFill.blip.rEmbed = rId
                replaced += 1
        for kind in non_picture:
            item = "%s（%s）" % (src.name, kind)
            if item not in non_picture_hits:
                non_picture_hits.append(item)
        names = picture_names(prs)
        for name in names:
            if name not in all_picture_names:
                all_picture_names.append(name)

        target = unique_path(out, src.name)
        if replaced == 0:
            # 这一份没有该形状：原样带过去，不清空产物目录，也不假装替换过
            shutil.copyfile(str(src), str(target))
            results.append({
                "input": str(src),
                "input_name": src.name,
                "output": str(target),
                "output_name": target.name,
                "replaced": 0,
                "image": image.name,
                "skipped": "没有名为 %s 的图片形状" % shape_name,
            })
            continue
        prs.save(str(target))
        results.append({
            "input": str(src),
            "input_name": src.name,
            "output": str(target),
            "output_name": target.name,
            "replaced": replaced,
            "image": image.name,
        })
        total += replaced

    if total == 0:
        if non_picture_hits:
            raise ToolError(
                "形状「%s」不是图片（%s）" % (shape_name, " / ".join(non_picture_hits[:6])),
                "replace-image 只能替换图片形状；文字替换请用 fill。",
            )
        raise ToolError(
            "没有名为「%s」的图片形状" % shape_name,
            ("这些文件里的图片形状名：" + " / ".join(all_picture_names[:12])) if all_picture_names
            else "这些文件里没有任何图片形状。",
        )

    skipped = sum(1 for r in results if r["replaced"] == 0)
    summary = "处理 %d 个文件，替换图片形状 %d 处（形状名：%s）" % (len(files), total, shape_name)
    if skipped:
        summary += "；%d 个文件没有该形状，已原样带过" % skipped
    return {
        "action": "replace-image",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "out": str(out),
        "files": results,
        "replaced": total,
        "skipped": skipped,
        "image": str(image),
        "summary": summary,
        "artifacts": [r["output"] for r in results],
    }


def slide_block(slide, index):
    lines = []
    title, title_id = "", None
    try:
        title_shape = slide.shapes.title
        if title_shape is not None and getattr(title_shape, "has_text_frame", False):
            title = title_shape.text_frame.text.strip()
            title_id = get_shape_id(title_shape)
    except Exception:
        title, title_id = "", None
    head = "=== 第 %d 页" % index
    if title:
        head += " · " + " ".join(title.split())
    lines.append(head + " ===")
    for shape in iter_shapes(slide.shapes):
        if getattr(shape, "has_table", False):
            rows = []
            for row in shape.table.rows:
                cells = [" ".join(c.text_frame.text.split()) for c in row.cells]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                lines.append("[表格 · %s]" % shape.name)
                lines.extend(rows)
            continue
        if not getattr(shape, "has_text_frame", False):
            continue
        if title_id is not None and get_shape_id(shape) == title_id:
            continue  # 标题已经在页头里
        text = shape.text_frame.text.strip()
        if not text:
            continue
        lines.append("[正文 · %s]" % shape.name)
        lines.extend(text.split("\n"))
    try:
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append("[备注]")
                lines.extend(notes.split("\n"))
    except Exception:
        pass  # 备注页异常不影响正文提取
    return "\n".join(lines)


def get_shape_id(shape):
    try:
        return shape.shape_id
    except Exception:
        return None


def action_extract_text(payload):
    files = collect_pptx(payload.get("input"))
    out = resolve_out(payload, payload.get("input"))
    out.mkdir(parents=True, exist_ok=True)
    results = []
    for src in files:
        prs = Presentation(str(src))
        blocks = [slide_block(slide, i) for i, slide in enumerate(prs.slides, 1)]
        text = "\n\n".join(blocks).strip() + "\n"
        target = unique_path(out, src.stem + ".txt")
        target.write_text(text, encoding="utf-8")
        results.append({
            "input": str(src),
            "input_name": src.name,
            "output": str(target),
            "output_name": target.name,
            "slides": len(prs.slides),
            "chars": len(text),
        })
    return {
        "action": "extract-text",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "out": str(out),
        "files": results,
        "summary": "提取 %d 个文件、共 %d 页文本" % (len(files), sum(r["slides"] for r in results)),
        "artifacts": [r["output"] for r in results],
    }


def action_info(payload):
    files = collect_pptx(payload.get("input"))
    results = []
    for src in files:
        prs = Presentation(str(src))
        shapes = sum(len(slide.shapes) for slide in prs.slides)
        pictures = 0
        for slide in prs.slides:
            pictures += sum(1 for shape in iter_shapes(slide.shapes) if is_picture(shape))
        results.append({
            "input": str(src),
            "input_name": src.name,
            "pages": len(prs.slides),
            "shapes": shapes,
            "pictures": pictures,
        })
    return {
        "action": "info",
        "input": str(Path(payload["input"]).expanduser().resolve()),
        "files": results,
        "pages": sum(r["pages"] for r in results),
        "summary": "共 %d 个文件、%d 页" % (len(files), sum(r["pages"] for r in results)),
        "artifacts": [],
    }


# ── demo ─────────────────────────────────────────────────────────

def write_solid_png(path, rgb, size=16):
    """手写一张纯色 PNG（不依赖 Pillow），给演示用的示例图片。"""
    width = height = size
    raw = b"".join(b"\x00" + bytes(rgb) * width for _ in range(height))

    def chunk(tag, data):
        return (struct.pack(">I", len(data)) + tag + data
                + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF))

    ihdr = struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0)
    png = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", ihdr)
           + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))
    Path(path).write_bytes(png)


def make_demo_deck(path, logo_path, client, amount, rows=None):
    """造一份带 {} 占位符 + 表格 + 图片形状的演示模板。"""
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # Title Only
    slide.shapes.title.text = "{{title}}"

    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(6.2), Inches(1.6))
    frame = box.text_frame
    frame.word_wrap = True
    frame.paragraphs[0].text = "客户：{{client}}"
    frame.add_paragraph().text = "金额：{{amount}}"

    table_rows = rows or [["项目", "内容"], ["{{client}}", "{{amount}}"]]
    table_shape = slide.shapes.add_table(
        len(table_rows), len(table_rows[0]), Inches(0.8), Inches(3.6), Inches(6.2), Inches(1.4)
    )
    for i, row in enumerate(table_rows):
        for j, value in enumerate(row):
            table_shape.table.cell(i, j).text = value

    picture = slide.shapes.add_picture(str(logo_path), Inches(7.2), Inches(0.8),
                                       Inches(2.2), Inches(2.2))
    picture.name = "Logo"
    prs.save(str(path))


def action_demo(payload):
    root = Path(tempfile.mkdtemp(prefix="dsh-pptx-batch-demo-"))
    inputs = root / "样例"
    inputs.mkdir(parents=True, exist_ok=True)
    old_logo = root / "logo-旧.png"
    new_logo = root / "logo-新.png"
    write_solid_png(old_logo, (200, 30, 30))
    write_solid_png(new_logo, (20, 120, 220))

    make_demo_deck(inputs / "简报A.pptx", old_logo, "{{client}}", "{{amount}}")
    make_demo_deck(inputs / "简报B.pptx", old_logo, "{{client}}", "{{amount}}",
                   rows=[["项目", "金额"], ["{{client}}", "{{amount}}"]])
    out = root / "输出"
    steps, artifacts = [], []

    def record(name, result):
        steps.append({
            "step": name,
            "action": result.get("action"),
            "summary": result.get("summary", ""),
            "out": result.get("out"),
        })
        artifacts.extend(result.get("artifacts", []))

    filled_dir = out / "fill"
    record("填占位符", action_fill({
        "input": str(inputs), "out": str(filled_dir),
        "set": ["title=" + DEMO_TITLE, "client=" + DEMO_CLIENT, "amount=" + DEMO_AMOUNT],
    }))
    record("换图片", action_replace_image({
        "input": str(filled_dir), "out": str(out / "replace-image"),
        "shape": "Logo", "image": str(new_logo),
    }))
    record("提文本", action_extract_text({
        "input": str(out / "replace-image"), "out": str(out / "extract-text"),
    }))
    record("看结构", action_info({"input": str(inputs)}))

    return {
        "action": "demo",
        "demoRoot": str(root),
        "inputRoot": str(inputs),
        "steps": steps,
        "summary": "演示完成：%d 个动作全部跑通，产物 %d 个" % (len(steps), len(artifacts)),
        "artifacts": artifacts,
    }


ACTIONS = {
    "fill": action_fill,
    "replace-image": action_replace_image,
    "extract-text": action_extract_text,
    "info": action_info,
    "demo": action_demo,
}


def main(argv):
    if len(argv) < 3:
        emit({"ok": False, "code": 2,
              "message": "内部调用参数不足：需要 <action> 和 JSON payload",
              "hint": "请通过 bin/cli.mjs 调用本脚本。"})
        return 2
    action, raw = argv[1], argv[2]
    try:
        payload = json.loads(raw) if raw else {}
    except json.JSONDecodeError as exc:
        emit({"ok": False, "code": 2, "message": "参数 JSON 解析失败：%s" % exc, "hint": ""})
        return 2
    fn = ACTIONS.get(action)
    if not fn:
        emit({"ok": False, "code": 2, "message": "未知动作：%s" % action,
              "hint": "可用动作：%s" % " / ".join(k for k in ACTIONS if k != "demo")})
        return 2
    try:
        result = fn(payload)
        result.setdefault("ok", True)
        emit(result)
        return 0
    except ToolError as exc:
        emit({"ok": False, "code": exc.code, "message": exc.message, "hint": exc.hint})
        return exc.code
    except PermissionError as exc:
        emit({"ok": False, "code": 1,
              "message": "没有权限读写文件：%s" % exc,
              "hint": "检查文件是否被 PowerPoint 占用、或输出目录是否有写权限。"})
        return 1
    except Exception as exc:  # noqa: BLE001 —— 兜底，保证永远输出 JSON
        traceback.print_exc(file=sys.stderr)
        emit({"ok": False, "code": 1, "message": "处理失败：%s" % exc,
              "hint": "文件可能已损坏或被加密；可先用 PowerPoint 打开另存一次再试。"})
        return 1


if __name__ == "__main__":
    sys.exit(main(sys.argv))
